"""Layer 3 — Render a Placement into a .pptx file via python-pptx.

This is the ONLY module that touches python-pptx.
No gradients, shadows, 3D, WordArt, animations.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

from microsoft_office.pptx_engine.design_system import (
    SLIDE_WIDTH,
    SLIDE_HEIGHT,
    TYPO,
    resolve_color,
    pt_to_emu,
)
from microsoft_office.pptx_engine.schemas import (
    ImageElement,
    Placement,
    RectElement,
    SlidePlacement,
    TextElement,
)
from microsoft_office.pptx_engine.validator import (
    assert_no_bold,
    validate_placement,
)


def _rgb(color: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*color)


def _apply_text_style(run, style_name: str) -> None:
    """Apply typography from design system to a run."""
    style = TYPO[style_name]
    run.font.name = style["font"]
    run.font.size = style["size"]
    run.font.bold = False  # Never bold
    assert_no_bold(False, context=f"style={style_name}")

    # Text color based on style
    if style_name == "headline":
        run.font.color.rgb = _rgb(resolve_color("mono_900"))
    elif style_name == "subhead":
        run.font.color.rgb = _rgb(resolve_color("mono_900"))
    elif style_name == "body":
        run.font.color.rgb = _rgb(resolve_color("mono_600"))
    elif style_name == "caption":
        run.font.color.rgb = _rgb(resolve_color("mono_400"))


def _add_textbox(slide, el: TextElement) -> None:
    """Add a textbox element to a slide."""
    left = Emu(pt_to_emu(el.region.left))
    top = Emu(pt_to_emu(el.region.top))
    width = Emu(pt_to_emu(el.region.width))
    height = Emu(pt_to_emu(el.region.height))

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    # Line spacing
    style = TYPO[el.style.value]
    p = tf.paragraphs[0]
    p.text = el.text
    p.space_after = Pt(0)
    p.line_spacing = style["line_spacing"]

    if p.runs:
        _apply_text_style(p.runs[0], el.style.value)
    else:
        # If no runs (empty or direct text), create one
        run = p.add_run()
        run.text = el.text
        p.text = ""  # Clear paragraph-level text
        p = tf.paragraphs[0]
        p.text = el.text
        if p.runs:
            _apply_text_style(p.runs[0], el.style.value)


def _add_rect(slide, el: RectElement) -> None:
    """Add a rectangle element to a slide."""
    from pptx.enum.shapes import MSO_SHAPE

    left = Emu(pt_to_emu(el.region.left))
    top = Emu(pt_to_emu(el.region.top))
    width = Emu(pt_to_emu(el.region.width))
    height = Emu(pt_to_emu(el.region.height))

    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(resolve_color(el.color.value))
    shape.line.fill.background()  # No border


def _add_image(slide, el: ImageElement) -> None:
    """Add an image element to a slide."""
    left = Emu(pt_to_emu(el.region.left))
    top = Emu(pt_to_emu(el.region.top))
    width = Emu(pt_to_emu(el.region.width))
    height = Emu(pt_to_emu(el.region.height))

    slide.shapes.add_picture(el.path, left, top, width, height)


def render(placement: Placement, output_path: str) -> str:
    """Render a Placement into a .pptx file.

    Args:
        placement: Validated Placement model.
        output_path: Where to save the .pptx file.

    Returns:
        Absolute path to the saved file.
    """
    # Validate before rendering
    errors = validate_placement(placement)
    if errors:
        raise ValueError(f"Placement validation failed:\n" + "\n".join(errors))

    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Use blank layout
    blank_layout = prs.slide_layouts[6]  # blank

    for slide_placement in placement.slides:
        slide = prs.slides.add_slide(blank_layout)

        # Set white background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(resolve_color("mono_000"))

        for el in slide_placement.elements:
            if el.type == "textbox":
                _add_textbox(slide, el)
            elif el.type == "rect":
                _add_rect(slide, el)
            elif el.type == "image":
                _add_image(slide, el)

    out = Path(output_path).resolve()
    prs.save(str(out))
    return str(out)
